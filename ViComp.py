import cv2
from ultralytics import YOLO
import time
from datetime import datetime as dt
import os
from openpyxl import Workbook, load_workbook

file_name = "XLlabWatch_.xlsx"

if not os.path.exists(file_name):
    workbook = Workbook()
    sheet = workbook.active
    sheet['A1'] = "time"
    sheet['B1'] = 'detections'
    sheet['C1'] = "% confident"
    workbook.save(filename=file_name)
else:
    workbook = load_workbook(filename=file_name)
    sheet = workbook.active

#Load YOLO model
model = YOLO('yolov8n.pt')

cap = cv2.VideoCapture(0, cv2.CAP_DSHOW)
cap.set(cv2.CAP_PROP_FRAME_WIDTH, 1000)
cap.set(cv2.CAP_PROP_FRAME_HEIGHT, 850)

if not cap.isOpened():
    print("error, camera could not be used")
    exit()

while True:
    work, frame = cap.read()
    if not work:
        break

    rgb_frame = cv2.cvtColor(frame, cv2.COLOR_BGR2RGB)
    results = model(rgb_frame, 640)

    fLog = []
    for result in results:
        boxes = result.boxes
        for box in boxes:
            x1, y1, x2, y2 = box.xyxy[0].int().tolist()
            confidence = box.conf[0].item()
            label = model.names[int(box.cls[0].item())]
            cv2.rectangle(frame, (x1, y1), (x2, y2), (100, 90, 225), 2)
            cv2.putText(frame, f'{label}: {confidence * 100:.2f}%', (x1, y1 - 10), cv2.FONT_HERSHEY_TRIPLEX, 0.8, (255, 255, 0), 1)
            
            cLog = [str(dt.now()), label, confidence * 100]
            fLog.append(cLog)

    #Append log to the sheet
    for log in fLog:
        sheet.append(log)

    cv2.imshow('YOLOv8 detection frame', frame)
    if cv2.waitKey(1) & 0xFF == ord('q'):
        break

    time.sleep(3)

workbook.save(filename=file_name)
cap.release()
cv2.destroyAllWindows()
